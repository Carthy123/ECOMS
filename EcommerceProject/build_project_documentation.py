from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE

OUTPUT = "Ecommerce_Project_Documentation.docx"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(table, color="D9D9D9"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    border = tbl_pr.first_child_found_in("w:tblBorders")
    if border is None:
        border = OxmlElement("w:tblBorders")
        tbl_pr.append(border)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = border.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            border.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell(cell, text, bold=False, color="000000"):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    borders(table)
    for index, header in enumerate(headers):
        set_cell(table.rows[0].cells[index], header, bold=True, color="FFFFFF")
        shade(table.rows[0].cells[index], "174A45")
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell(cells[index], str(value))
            if row_index % 2 == 1:
                shade(cells[index], "F3F7F5")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.add_run(text)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].font.color.rgb = RGBColor(35, 55, 51)
styles["Normal"].paragraph_format.space_after = Pt(7)
for style_name, size in (("Title", 25), ("Heading 1", 17), ("Heading 2", 12.5)):
    style = styles[style_name]
    style.font.name = "Aptos Display" if style_name == "Title" else "Aptos"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(14 if style_name != "Title" else 0)
    style.paragraph_format.space_after = Pt(7)

title = doc.add_paragraph(style="Title")
title.add_run("Ecommerce Project Documentation")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(18)
run = subtitle.add_run("Merchant storefront and ecommerce management workspace")
run.italic = True
run.font.size = Pt(11)
run.font.color.rgb = RGBColor(15, 118, 110)

doc.add_heading("Project purpose", level=1)
doc.add_paragraph(
    "EcommerceProject is a full-stack ecommerce store for merchants who need to manage a product catalogue, inventory, orders, customers, and storefront presentation from one workspace. The current implementation provides a working local development experience with a React frontend, an ASP.NET Core API, and a SQLite development database."
)
doc.add_paragraph(
    "This document explains what has been built, how the main user journeys work, how to run the project, and which production features remain to be connected."
)

doc.add_heading("Current feature set", level=1)
add_table(doc, ["Area", "Implemented capability"], [
    ("Authentication", "Merchant registration, login, JWT session handling, sign out, and demo mode."),
    ("Store setup", "Create a store and edit its name, slug, description, email, and phone number."),
    ("Catalogue", "Create products and categories, upload product images, edit product details, and adjust stock."),
    ("Product images", "Uploaded images are supported, with remote product image fallbacks for the starter catalogue."),
    ("Storefront", "Responsive product grid, product cards, image presentation, add-to-cart actions, and animated interactions."),
    ("Cart and checkout", "Quantity controls, totals, customer details, order submission, stock validation, and confirmation."),
    ("Orders", "Merchant order list with customer details, line items, delivery address, totals, and status updates."),
    ("Customers", "Customer summaries derived from orders, including order count and total spend."),
    ("Product lifecycle", "Products can be archived without deleting their historical order records."),
])

doc.add_heading("How to use the application", level=1)
doc.add_heading("Merchant workflow", level=2)
for text in [
    "Open the dashboard and sign in or register a merchant account.",
    "Create a store, then open Store settings to maintain the store identity and contact details.",
    "Use Add product to create a product with name, SKU, price, stock, category, and optional image.",
    "Use the three-dot action in the product table to edit a product or archive it.",
    "Use the stock controls to make small inventory adjustments.",
    "Select View storefront to preview the customer-facing shop.",
]:
    add_bullet(doc, text)
doc.add_heading("Customer workflow", level=2)
for text in [
    "Browse active products and view their images, names, SKUs, and prices.",
    "Add products to the cart. Repeated additions increase the quantity of an existing cart line.",
    "Adjust quantities or remove items from the cart.",
    "Complete checkout with customer name, email, and delivery address.",
    "The API validates stock, creates the order, deducts inventory, and returns an order reference.",
]:
    add_bullet(doc, text)

doc.add_heading("Technical architecture", level=1)
add_table(doc, ["Layer", "Technology and responsibility"], [
    ("Frontend", "React, TypeScript, Vite, and CSS. The dashboard and storefront live in src/EcommerceProject.Web."),
    ("API", "ASP.NET Core .NET 8 REST API in src/EcommerceProject.Api."),
    ("Authentication", "ASP.NET Core Identity with JWT bearer authentication and merchant authorization."),
    ("Data access", "Entity Framework Core with SQLite for local development and SQL Server configuration for production."),
    ("Local API", "http://127.0.0.1:5087"),
    ("Local frontend", "http://127.0.0.1:5173"),
])

doc.add_heading("Important API routes", level=1)
add_table(doc, ["Method", "Route", "Purpose"], [
    ("POST", "/api/auth/register", "Register a merchant account."),
    ("POST", "/api/auth/login", "Authenticate and receive a JWT."),
    ("GET", "/api/stores", "List stores available to the merchant."),
    ("PUT", "/api/stores/{id}", "Update store settings."),
    ("GET", "/api/stores/{storeId}/products", "List active products for a store."),
    ("POST", "/api/stores/{storeId}/products", "Create a product."),
    ("PUT", "/api/products/{id}", "Edit or archive a product."),
    ("POST", "/api/products/{id}/image", "Upload a product image."),
    ("POST", "/api/orders", "Create an order and deduct stock."),
    ("GET", "/api/orders/store/{storeId}", "List merchant orders."),
    ("PATCH", "/api/orders/{id}/status", "Update an order status."),
])

doc.add_heading("Local setup", level=1)
doc.add_paragraph("Run the API and frontend from the EcommerceProject folder. The commands below assume the required .NET SDK and Node.js dependencies are installed.")
for text in [
    "API: dotnet run --project src/EcommerceProject.Api",
    "Frontend: cd src/EcommerceProject.Web, then npm run dev",
    "API health check: Invoke-RestMethod http://127.0.0.1:5087/health",
    "Frontend build check: npm run build from src/EcommerceProject.Web",
]:
    add_bullet(doc, text)

doc.add_heading("Validation and verification", level=1)
for text in [
    "The frontend production build has been run successfully after the latest cart, image, settings, product editing, and storefront layout changes.",
    "The API health endpoint has returned status ok during development verification.",
    "Order creation validates that the store exists, requested products are active, requested quantities are positive, and stock is available.",
    "The storefront filters archived products out of the customer-facing product list.",
]:
    add_bullet(doc, text)

doc.add_heading("Remaining production work", level=1)
doc.add_paragraph("The project is functional as a local ecommerce prototype. The following items should be completed before public release:")
for text in [
    "Connect a real payment provider and add payment confirmation handling.",
    "Add production email notifications for order confirmation and fulfilment updates.",
    "Add complete archived-product browsing and restore controls in the merchant catalogue.",
    "Add richer category and stock filters to the product table.",
    "Configure production secrets, SQL Server connection settings, HTTPS, deployment, backups, and monitoring.",
    "Review image licensing and replace development-only remote image URLs with approved production assets.",
]:
    add_bullet(doc, text)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer.add_run("EcommerceProject | Project documentation")
footer_run.font.size = Pt(8)
footer_run.font.color.rgb = RGBColor(115, 131, 126)

doc.save(OUTPUT)
print(OUTPUT)
